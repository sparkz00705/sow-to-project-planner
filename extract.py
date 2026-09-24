
from __future__ import annotations

from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_document(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts).strip()

    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    parts.append(" | ".join(vals))
        return "\n".join(parts).strip()

    return data.decode("utf-8", errors="ignore").strip()
