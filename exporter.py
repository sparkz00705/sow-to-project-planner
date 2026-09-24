
from __future__ import annotations

from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)


def _sheet(wb, name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            widths[cell.column_letter] = min(max(widths.get(cell.column_letter, 0), len(str(cell.value or "")) + 2), 45)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    return ws


def build_excel_workbook(plan: dict) -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    s = plan.get("summary", {})
    _sheet(wb, "Project Summary", ["Field", "Value"], [
        ["Project Name", s.get("project_name", "")],
        ["Project Type", s.get("project_type", "")],
        ["Description", s.get("description", "")],
        ["Confidence", s.get("confidence", "")],
        ["Engine", plan.get("metadata", {}).get("engine", "")],
        ["Traceability Coverage", f"{plan.get('metadata', {}).get('traceability_coverage_percent', 0)}%"],
    ])

    _sheet(wb, "SOW Items",
           ["SOW ID","Section","Statement","Type","Executable","Priority"],
           [[x.get("sow_id"),x.get("section"),x.get("statement"),x.get("type"),x.get("executable"),x.get("priority")] for x in plan.get("sow_items",[])])

    _sheet(wb, "WBS", ["WBS ID","Phase","Objective"], [
        [x.get("wbs_id"),x.get("phase"),x.get("objective")] for x in plan.get("wbs",[])
    ])

    _sheet(wb, "Project Plan",
           ["WBS","Activity ID","Activity","Phase","Duration (days)","Dependencies","Owner","Deliverable","Milestone","Source SOW IDs","Planning Basis"],
           [[x.get("wbs_id"),x.get("activity_id"),x.get("activity_name"),x.get("phase"),x.get("duration_days"),
             ", ".join(x.get("dependency_ids",[])),x.get("owner_role"),x.get("deliverable"),x.get("milestone"),
             ", ".join(x.get("source_sow_ids",[])),x.get("planning_note")] for x in plan.get("activities",[])])

    _sheet(wb, "Milestones", ["ID","Name","Target","Source","Source SOW IDs"], [
        [x.get("milestone_id"),x.get("name"),x.get("target"),x.get("source"),", ".join(x.get("source_sow_ids",[]))]
        for x in plan.get("milestones",[])
    ])

    _sheet(wb, "Traceability", ["SOW ID","Item Type","Executable","Activity IDs","Milestone IDs","Status","Reason"], [
        [x.get("sow_id"),x.get("item_type"),x.get("executable"),", ".join(x.get("activity_ids",[])),
         ", ".join(x.get("milestone_ids",[])),x.get("status"),x.get("reason")] for x in plan.get("traceability",[])
    ])

    _sheet(wb, "Gaps", ["ID","Category","Description","Severity","Recommendation","Source SOW IDs"], [
        [x.get("gap_id"),x.get("category"),x.get("description"),x.get("severity"),x.get("recommendation"),", ".join(x.get("source_sow_ids",[]))]
        for x in plan.get("gaps",[])
    ])

    _sheet(wb, "Risks", ["ID","Risk","Impact","Probability","Mitigation","Source SOW IDs"], [
        [x.get("risk_id"),x.get("risk"),x.get("impact"),x.get("probability"),x.get("mitigation"),", ".join(x.get("source_sow_ids",[]))]
        for x in plan.get("risks",[])
    ])

    _sheet(wb, "Assumptions", ["ID","Assumption","Basis","Status","Source SOW IDs"], [
        [x.get("assumption_id"),x.get("assumption"),x.get("basis"),x.get("status"),", ".join(x.get("source_sow_ids",[]))]
        for x in plan.get("assumptions",[])
    ])

    _sheet(wb, "Constraints", ["ID","Constraint","Source SOW IDs"], [
        [x.get("constraint_id"),x.get("constraint"),", ".join(x.get("source_sow_ids",[]))]
        for x in plan.get("constraints",[])
    ])

    _sheet(wb, "Raw JSON", ["Section","JSON"], [
        ["Plan", json.dumps(plan, ensure_ascii=False)]
    ])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
