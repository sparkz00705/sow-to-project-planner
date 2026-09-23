from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _frame(items):
    if not items:
        return pd.DataFrame([{"Info": "No records"}])
    return pd.DataFrame(items)


def build_excel_workbook(plan: dict) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary = plan.get("summary", {})
        pd.DataFrame(
            [
                ["Project Name", summary.get("project_name", "")],
                ["Project Type", summary.get("project_type", "")],
                ["Description", summary.get("description", "")],
                ["Confidence", summary.get("confidence", "")],
            ],
            columns=["Field", "Value"],
        ).to_excel(writer, sheet_name="Project Summary", index=False)
        _frame(plan.get("sow_items", [])).to_excel(writer, sheet_name="SOW Items", index=False)
        _frame(plan.get("wbs", [])).to_excel(writer, sheet_name="WBS", index=False)
        _frame(plan.get("activities", [])).to_excel(writer, sheet_name="Project Plan", index=False)
        _frame(plan.get("milestones", [])).to_excel(writer, sheet_name="Milestones", index=False)
        _frame(plan.get("traceability", [])).to_excel(writer, sheet_name="Traceability", index=False)
        _frame(plan.get("gaps", [])).to_excel(writer, sheet_name="Gaps", index=False)
        _frame(plan.get("risks", [])).to_excel(writer, sheet_name="Risks", index=False)
        pd.DataFrame({"Assumptions": plan.get("assumptions", []) or [""]}).to_excel(
            writer, sheet_name="Assumptions", index=False
        )
        pd.DataFrame({"Constraints": plan.get("constraints", []) or [""]}).to_excel(
            writer, sheet_name="Constraints", index=False
        )

    output.seek(0)
    wb = load_workbook(output)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        for col_idx, column_cells in enumerate(ws.columns, 1):
            max_len = min(max((len(str(c.value)) if c.value is not None else 0) for c in column_cells) + 2, 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = max(12, max_len)
    final = BytesIO()
    wb.save(final)
    final.seek(0)
    return final.getvalue()
